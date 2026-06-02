"""
Converts an Excel workbook (.xlsx) to CSV files, one per sheet.
Output goes to data/processed/<sheet_name>.csv

Usage:
    python scripts/prepare_data.py data/your-file.xlsx
"""

import sys
import os
import csv

try:
    import openpyxl
except ImportError:
    print("openpyxl is required. Install it with: pip install openpyxl")
    sys.exit(1)


def sanitize_name(name: str) -> str:
    return "".join(c if c.isalnum() or c in "-_" else "_" for c in name).strip("_")


def convert(excel_path: str) -> None:
    if not os.path.isfile(excel_path):
        print(f"File not found: {excel_path}")
        sys.exit(1)

    out_dir = os.path.join("data", "processed")
    os.makedirs(out_dir, exist_ok=True)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    print(f"Loaded: {excel_path} ({len(wb.sheetnames)} sheet(s) found)\n")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        # Skip completely empty sheets
        if not any(any(cell is not None for cell in row) for row in rows):
            print(f"  Skipping '{sheet_name}' — sheet is empty")
            continue

        safe_name = sanitize_name(sheet_name)
        out_path = os.path.join(out_dir, f"{safe_name}.csv")

        with open(out_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for row in rows:
                writer.writerow(["" if cell is None else cell for cell in row])

        data_rows = sum(1 for r in rows if any(c is not None for c in r))
        print(f"  '{sheet_name}' -> {out_path}  ({data_rows} non-empty rows)")

    print(f"\nDone. CSV files are in: {out_dir}/")
    print("Reference them in Claude Code chat with @data/processed/<filename>.csv")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python scripts/prepare_data.py <path-to-excel-file>")
        sys.exit(1)
    convert(sys.argv[1])
